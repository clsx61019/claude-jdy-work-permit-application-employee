"""
一次性工具：把简道云的三份授权书模板转换成本项目的模板。

- Authorization for IMO（移民局授权书，匈英双语，公司自己的格式）
- NAV Authorization（税务局授权书，改自 NAV「eseti meghatalmazásminta」推荐样本）
- TAJ Authorization（TAJ 卡授权书，改自 NEAK NYT.159 表格）

${标签#_widget_xxx} 和写死的被授权人、证人换成 {{名字}}（值由 service/forms/*_authorization.py 计算），
并做 docs/rules/授权书.md 里的 A 类修正。

用法: python tools/build_authorizations.py
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from service.docx_text import para_text, paragraphs, replace_in_paragraph  # noqa: E402

W = r"\$\{[^}]*?#_widget_"  # 占位符开头


def fix(doc, old, new, regex=False, expect=1):
    pattern = old if regex else re.escape(old)
    n = sum(replace_in_paragraph(p, pattern, lambda m: new) for p in paragraphs(doc))
    if n != expect:
        raise ValueError(f"「{old[:50]}…」应命中 {expect} 处，实际 {n} 处")


# 原模板里写死的被授权人、证人是真人信息，入库前已换成 [NÉV]、[CÍM] 之类的标记

def check_done(doc, leftovers=()):
    left = [para_text(p) for p in paragraphs(doc) if "${" in para_text(p)]
    for word in leftovers:  # 标记必须全部换掉
        left += [para_text(p) for p in paragraphs(doc) if word in para_text(p)]
    if left:
        raise ValueError(f"还有没换掉的内容：{left}")


def build_imo():
    doc = Document(ROOT / "templates/jdy/imo_authorization_jdy_original.docx")
    fix(doc, W + r"1728475331388\}", "{{name}}", regex=True, expect=2)
    fix(doc, W + r"1773645284570\}", "{{birth_date_en}}", regex=True)
    fix(doc, W + r"1773645284571\}", "{{birth_date_hu}}", regex=True)
    fix(doc, W + r"1728475331413\}", "{{mother_name}}", regex=True, expect=2)
    fix(doc, W + r"1732625466696\}", "{{auth_name}}", regex=True, expect=2)
    # A8：被授权人出生日期原来用 3 个辅助字段拼成 1981.02.14，英文里也是匈牙利格式
    fix(doc, W + r"1732625466821\}\." + W + r"1732625466822\}\." + W + r"1732625466823\}, ",
        "{{auth_birth_date_en}}, ", regex=True)
    fix(doc, W + r"1732625466821\}\." + W + r"1732625466822\}\." + W + r"1732625466823\}\., ",
        "{{auth_birth_date_hu}}, ", regex=True)
    fix(doc, W + r"1732625466698\}", "{{auth_mother}}", regex=True, expect=2)
    fix(doc, W + r"1732625466780\}", "{{auth_position}}", regex=True, expect=2)
    fix(doc, W + r"1732625466781\}", "{{auth_company}}", regex=True, expect=2)
    # A7
    # 匈语冠词按公司名首字母：a BorsodChem / az Eastern…（代码算）
    fix(doc, "a {{auth_company}} {{auth_position}} munkatársát", "{{auth_company_hu}} {{auth_position}} munkatársát")
    fix(doc, "Országos ldegenrendészeti Főigazgatóság", "Országos Idegenrendészeti Főigazgatóság")
    # B2：签字地点按「递交地点」；两栏都填签字日期
    fix(doc, "Beijing, ………………………", "{{place_en}}, {{sign_date_en}}")
    fix(doc, "Kazincbarcika, " + "${辅助_签字日期1（匈牙利语）#_widget_1773645284868}",
        "{{place_hu}}, {{sign_date_hu}}")
    for i, (name, idno, addr) in enumerate([("943", "944", "945"), ("946", "947", "948")], 1):
        fix(doc, W + rf"1733307002{name}\}}", f"{{{{w{i}_name}}}}", regex=True)
        fix(doc, W + rf"1733307002{idno}\}}", f"{{{{w{i}_id}}}}", regex=True)
        fix(doc, W + rf"1733307002{addr}\}}", f"{{{{w{i}_address}}}}", regex=True)
    check_done(doc)
    doc.save(ROOT / "templates/hu/imo_authorization.docx")


def build_nav():
    doc = Document(ROOT / "templates/jdy/nav_authorization_jdy_original.docx")
    # A4：官方样本的脚注正文被删了，正文里只剩脚注号（第一个还是撇号）
    fix(doc, "Meghatalmazó adatai'", "Meghatalmazó adatai")
    fix(doc, r"Meghatalmazott adatai 2 $", "Meghatalmazott adatai", regex=True)
    fix(doc, r"eljárjon: 3 $", "eljárjon:", regex=True)
    fix(doc, r"tanúk előtt: 4 $", "tanúk előtt:", regex=True)
    # A6
    fix(doc, r"^,hogy a Nemzeti", ", hogy a Nemzeti", regex=True)
    # 委托人
    fix(doc, W + r"1728475331388\}", "{{name}}", regex=True)
    fix(doc, W + r"1756727041750\}", "{{tax_no}}", regex=True)
    # A5：出生地加国家
    fix(doc, W + r"1732527942838\}, " + W + r"1773645284571\}",
        "{{birth_place}}, {{birth_country_hu}}, {{birth_date_hu}}", regex=True)
    fix(doc, W + r"1728475331413\}", "{{mother_name}}", regex=True)
    # B3：委托人通信地址保持 BC 地址（模板里写死，不动）
    # B1：被授权人原来写死，改为取简道云「选择被授权人」
    fix(doc, "Családi és utóneve: [NÉV]", "Családi és utóneve: {{auth_name}}")
    fix(doc, "Adóazonosító jele: [ADÓAZONOSÍTÓ]", "Adóazonosító jele: {{auth_tax_no}}")
    fix(doc, "Születési helye, ideje: [HELY], [DÁTUM]", "Születési helye, ideje: {{auth_birth_place}}, {{auth_birth_date_hu}}")
    fix(doc, "Anyja neve: [ANYJA NEVE]", "Anyja neve: {{auth_mother}}")
    fix(doc, "székhelye: [CÍM]", "székhelye: {{auth_address}}")
    # 授权事项原来是黄色高亮，打印出来有底色
    for par in paragraphs(doc):
        if para_text(par) == "Jövedelem igazolás az előző évről":
            for h in par.iter(qn("w:highlight")):
                h.getparent().remove(h)
    # B5：签字日期
    fix(doc, "Kelt: Kazincbarcika,…………………….", "Kelt: Kazincbarcika, {{sign_date_hu}}")
    for i, (name, addr, mother) in enumerate([("1733307002943", "1733307002945", "1756728816075"),
                                              ("1733307002946", "1733307002948", "1756728816076")], 1):
        fix(doc, W + name + r"\}", f"{{{{w{i}_name}}}}", regex=True)
        fix(doc, W + addr + r"\}", f"{{{{w{i}_address}}}}", regex=True)
        fix(doc, W + mother + r"\}", f"{{{{w{i}_mother}}}}", regex=True)
    check_done(doc, ["["])
    doc.save(ROOT / "templates/hu/nav_authorization.docx")


def build_taj():
    doc = Document(ROOT / "templates/jdy/taj_authorization_jdy_original.docx")
    fix(doc, W + r"1728475331388\}", "{{name}}", regex=True, expect=2)
    fix(doc, W + r"1732527942838\}, " + W + r"1732527942947\}", "{{birth_place}}, {{birth_country_hu}}", regex=True)
    fix(doc, W + r"1728562133707\}\. év " + W + r"1728562133738\}\. hó " + W + r"1728562133766\}\. nap",
        "{{birth_date}}", regex=True)
    fix(doc, W + r"1728475331413\}", "{{mother_name}}", regex=True)
    # A1：值和标签分开；街道名称后原来又印了「utca」（Ipari útutca）
    fix(doc, W + r"1728549508506\} irsz\.", "{{hu_zip}} irsz.", regex=True)
    fix(doc, W + r"1728549508516\}helység", "{{hu_city}} helység", regex=True)
    # 门牌号那格很窄，加上楼层、室会折成三行：街道和门牌号写在同一格（Ipari út 9-11. hsz, 2. em., 3. ajtó）
    fix(doc, W + r"1728549508519\} " + W + r"1728549508520\}utca", "{{hu_street}}", regex=True)
    fix(doc, W + r"1728549508521\} hsz", "", regex=True)
    # A2：委托人是护照号
    fix(doc, "személyi ig. szám:", "útlevél szám:", expect=2)  # 先把两处都改掉，被授权人那处下面改回
    fix(doc, W + r"1728480011304\}", "{{passport_no}}", regex=True)
    # A3：签字日期用「签字日期」，按官方写法
    fix(doc, W + r"1733824808671\}\. év " + W + r"1733824808672\}\. hó " + W + r"1733824808673\}\. nap ",
        "{{sign_date}}", regex=True)
    # B1：被授权人和证人原来写死，改为取简道云。住址是一个文本字段，写在第一行，其余几行留空
    fix(doc, "[NÉV]-t", "{{auth_name}}-t")
    fix(doc, r"^\[NÉV\]$", "{{auth_name}}", regex=True)
    fix(doc, "[HELY]", "{{auth_birth_place}}")
    fix(doc, "[DÁTUM]", "{{auth_birth_date}}")
    fix(doc, "[ANYJA NEVE]", "{{auth_mother}}")
    fix(doc, "[IRSZ] irsz.", "{{auth_address}}")
    for old in ("[HELYSÉG]helység", "[UTCA]utca", "[SZÁM]. hsz", "[T1 HELYSÉG]helység", "[T2 HELYSÉG]helység",
                "[T1 UTCA]utca", "[T1 SZÁM]. hsz", "[T2 UTCA]utca", "[T2 SZÁM]. hsz"):
        fix(doc, old, "")
    fix(doc, "[SZIG. SZÁM]", "{{auth_id_no}}")
    fix(doc, "[TANÚ 1]", "{{w1_name}}")
    fix(doc, "[TANÚ 2]", "{{w2_name}}")
    fix(doc, "[T1 IRSZ] irsz.", "{{w1_address}}")
    fix(doc, "[T2 IRSZ] irsz.", "{{w2_address}}")
    # 被授权人一栏仍是身份证号：两处「útlevél szám:」里第二处（被授权人栏）改回
    labels = [p for p in paragraphs(doc) if para_text(p) == "útlevél szám:"]
    replace_in_paragraph(labels[1], re.escape("útlevél szám:"), lambda m: "személyi ig. szám:")
    # 证人住址原来分三行，现在是一个字段：删掉空出来的两行
    witness = [t for t in doc.tables if "1. tanú:" in para_text(t._tbl)][0]
    for tr in [r._tr for r in witness.rows if not para_text(r._tr).strip()]:
        tr.getparent().remove(tr)
    check_done(doc, ["["])
    doc.save(ROOT / "templates/hu/taj_authorization.docx")


if __name__ == "__main__":
    build_imo()
    build_nav()
    build_taj()
    print("已生成 templates/hu/imo_authorization.docx、nav_authorization.docx、taj_authorization.docx")
