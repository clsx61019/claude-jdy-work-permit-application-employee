"""
一次性工具：把简道云的 TAJ 两份表格和 GDPR 声明转换成本项目的模板（规则见 docs/rules/TAJ申请和GDPR声明.md）。

- TAJ Application Form → taj_application.docx（NEAK NYT.53 Igénylőlap，TAJ 申请表）
- TAJ Card Application → taj_order.docx（NEAK NYT.52 Megrendelő，雇主的 TAJ 订购单）
- GDPR → gdpr_declarations.docx（告知确认 + 数据处理同意书，英语）

TAJ 原模板和 NEAK 官方 NYT.52.G / NYT.53.G（2020）逐字对过，只有 NYT.53 脚注的法条不同，改成官方写法。
原模板里写死的递交人（姓名和家庭住址）入库前已换成 [NÉV]、[TELEPÜLÉS] 之类的标记。

复选框改成有意义的名字，由 service/forms/*.py 用 doc.check(名字) 勾选。

用法: python tools/build_taj_gdpr.py
"""
import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from service.docx_text import para_text, paragraphs, replace_in_paragraph  # noqa: E402

W = r"\$\{[^}]*?#_widget_"  # 占位符开头
OUT = ROOT / "templates/hu"


def fix(doc, old, new, regex=False, expect=1):
    pattern = old if regex else re.escape(old)
    n = sum(replace_in_paragraph(p, pattern, lambda m: new) for p in paragraphs(doc))
    if n != expect:
        raise ValueError(f"「{old[:50]}…」应命中 {expect} 处，实际 {n} 处")


def widget(doc, wid, token, expect=1):
    fix(doc, W + wid + r"\}", token, regex=True, expect=expect)


def check_done(doc, leftovers=()):
    left = [para_text(p) for p in paragraphs(doc) if "${" in para_text(p)]
    for word in leftovers:
        left += [para_text(p) for p in paragraphs(doc) if word in para_text(p)]
    if left:
        raise ValueError(f"还有没换掉的内容：{left}")


def name_checkboxes(doc, names):
    """按文档顺序给复选框起名字（原来是 Jelölő2…8，而且有重名），并全部设为不勾"""
    boxes = [ff for ff in doc.element.body.iter(qn("w:ffData")) if ff.find(qn("w:checkBox")) is not None]
    if len(boxes) != len(names):
        raise ValueError(f"复选框应有 {len(names)} 个，实际 {len(boxes)} 个")
    for ff, name in zip(boxes, names):
        old = ff.find(qn("w:name")).get(qn("w:val"))
        ff.find(qn("w:name")).set(qn("w:val"), name)
        # 紧跟着的书签同名，一起改
        run = ff.getparent().getparent()
        bm = run.getnext()
        if bm is not None and bm.tag == qn("w:bookmarkStart") and bm.get(qn("w:name")) == old:
            bm.set(qn("w:name"), name)
        box = ff.find(qn("w:checkBox"))
        for c in box.findall(qn("w:checked")):
            box.remove(c)


def remove_highlight(doc):
    """原模板里性别、日期等处有黄色高亮，打印出来有底色"""
    for h in list(doc.element.body.iter(qn("w:highlight"))):
        h.getparent().remove(h)


def rows(doc, table):
    return list(doc.element.body.iter(qn("w:tbl")))[table].findall(qn("w:tr"))


def set_cell(tc, text, like):
    """单元格写成 text：有文字的换掉，空的照抄 like 段落里第一个 run 的格式写进去"""
    p = tc.find(qn("w:p"))
    if para_text(p):
        assert replace_in_paragraph(p, "^.*$", lambda m: text) == 1
        return
    run = copy.deepcopy(next(like.iter(qn("w:r"))))
    for t in run.findall(qn("w:t"))[1:]:
        run.remove(t)
    run.find(qn("w:t")).text = text
    p.append(run)


def date_cells(tr, y, m, d):
    """「… | 空 | év | 空 | hó | 空 | nap」这种行：在 év、hó、nap 前面的格子里填年、月、日"""
    tcs = tr.findall(qn("w:tc"))
    texts = [para_text(tc).strip() for tc in tcs]
    for label, token in (("év", y), ("hó", m), ("nap", d)):
        i = texts.index(label)
        set_cell(tcs[i - 1], token, tcs[i].find(qn("w:p")))


def person(doc, birth_day_widget):
    """两份 TAJ 表格共用的员工部分"""
    widget(doc, "1728475331386", "{{surname}}")
    widget(doc, "1728475331387", "{{given_name}}")
    widget(doc, "1728562133650", "{{birth_surname}}")
    widget(doc, "1728562133651", "{{birth_given_name}}")
    widget(doc, "1728475331411", "{{mother_surname}}")
    widget(doc, "1728475331412", "{{mother_given_name}}")
    widget(doc, r"1732527942838\}, " + W + "1732527942947", "{{birth_place}}, {{birth_country_hu}}")
    widget(doc, "1728562133707", "{{birth_y}}")
    # A1：订购单原来「日」用了「月」的字段
    if birth_day_widget == "1728562133738":
        fix(doc, r"^" + W + r"1728562133738\}$", "{{birth_m}}", regex=True, expect=2)
        texts = [p for p in paragraphs(doc) if para_text(p) == "{{birth_m}}"]
        replace_in_paragraph(texts[1], re.escape("{{birth_m}}"), lambda m: "{{birth_d}}")
    else:
        widget(doc, "1728562133738", "{{birth_m}}")
        widget(doc, birth_day_widget, "{{birth_d}}")
    widget(doc, "1732550878139", "{{nationality_hu}}")
    # 居住地、居留地都填匈牙利住址（B2）；楼层、室按字段（A3）
    widget(doc, "1728549508506", "{{hu_zip}}", expect=2)
    widget(doc, "1728549508516", "{{hu_city}}", expect=2)
    fix(doc, r"^" + W + r"1728549508519\} " + W + r"1728549508520\} ?$", "{{hu_street}}", regex=True, expect=2)
    widget(doc, "1728549508521", "{{hu_house_no}}", expect=2)


def floor_cells(doc, table, row_indexes, token):
    for i in row_indexes:
        tcs = rows(doc, table)[i].findall(qn("w:tc"))
        assert para_text(tcs[-2]) == "emelet, ajtó" and para_text(tcs[-1]) == "-", para_text(rows(doc, table)[i])
        set_cell(tcs[-1], token, None)


def build_application():
    doc = Document(ROOT / "templates/jdy/taj_application_jdy_original.docx")
    name_checkboxes(doc, ["first_issue", "lost", "name_change", "expired", "male", "female"])
    person(doc, "1728562133766")
    floor_cells(doc, 1, (9, 11), "{{hu_floor_door}}")
    # A2：递交人原来写死，改为被授权人；住址是一个字段（待办 J15），整行合成一格
    fix(doc, "[NÉV]", "{{auth_name}}")
    trs = rows(doc, 1)
    addr, street = trs[13], trs[14]
    tcs = addr.findall(qn("w:tc"))
    assert para_text(tcs[0]) == "Lakóhelye, székhelye"
    keep, span, width = tcs[-1], 0, 0
    for tc in tcs[1:]:
        pr = tc.find(qn("w:tcPr"))
        gs = pr.find(qn("w:gridSpan"))
        span += int(gs.get(qn("w:val"))) if gs is not None else 1
        width += int(pr.find(qn("w:tcW")).get(qn("w:w")))
        if tc is not keep:
            addr.remove(tc)
    pr = keep.find(qn("w:tcPr"))
    pr.find(qn("w:gridSpan")).set(qn("w:val"), str(span))
    pr.find(qn("w:tcW")).set(qn("w:w"), str(width))
    set_cell(keep, "{{auth_address}}", None)
    assert "[UTCA]" in para_text(street)
    street.getparent().remove(street)
    # 落款日期：生成当天（B5）
    date_cells(rows(doc, 1)[15], "{{today_y}}", "{{today_m}}", "{{today_d}}")
    # 脚注法条按官方 NYT.53.G（2020）
    fix(doc, "Első TAJ kiadásához a Tbj. 16.§ alá tartozó", "Első TAJ kiadásához az új Tbj. 22.§ alá tartozó")
    remove_highlight(doc)
    check_done(doc, ["["])
    doc.save(OUT / "taj_application.docx")


def build_order():
    doc = Document(ROOT / "templates/jdy/taj_order_jdy_original.docx")
    name_checkboxes(doc, ["employee", "member", "egt_yes", "egt_no", "commuter_yes", "commuter_no",
                          "simplified_yes", "simplified_no", "male", "female"])
    # A5：雇主信息取 config/employers.json
    widget(doc, "1732604018708", "{{employer}}")
    widget(doc, "1732695761783", "{{employer_zip}}")
    widget(doc, "1732695761784", "{{employer_city}}")
    widget(doc, r"1732695761785\} " + W + "1732695761786", "{{employer_street}}")
    widget(doc, "1732695761787", "{{employer_house_no}}")
    widget(doc, "1732604018791", "{{employer_tax_no}}")
    person(doc, "1728562133738")
    floor_cells(doc, 1, (9, 11), "{{hu_floor_door}}")
    # B3：在现雇主的入职日期 = 劳动关系实际起始日期
    date_cells(rows(doc, 1)[12], "{{start_y}}", "{{start_m}}", "{{start_d}}")
    # B5：落款日期 = 生成当天
    fix(doc, "Kazincbarcika,     év      hó      nap ", "Kazincbarcika, {{today}}")
    # B6：签字栏下印雇主代表
    sig = next(p for p in paragraphs(doc) if para_text(p) == "P.Hfoglalkoztató")
    rep = copy.deepcopy(sig)
    sig.addnext(rep)
    replace_in_paragraph(rep, "P.H", lambda m: "")
    replace_in_paragraph(rep, "foglalkoztató", lambda m: "{{representative}}")
    remove_highlight(doc)
    check_done(doc)
    doc.save(OUT / "taj_order.docx")


def build_gdpr():
    doc = Document(ROOT / "templates/jdy/gdpr_declarations_jdy_original.docx")
    widget(doc, "1728475331388", "{{name}}", expect=2)
    # A6：地址和劳动合同同一种写法
    fix(doc, r"^" + W + r"1732527943145\}, .*" + W + r"1743597826298\}$", "{{home_address}}", regex=True, expect=2)
    widget(doc, "1773645284570", "{{birth_date}}", expect=2)
    # B5：签字地点按递交地点、日期 = 签字日期（和员工签的其他声明一样）
    fix(doc, r"^Kazincbarcika, …+$", "{{place}}, {{sign_date}}", regex=True, expect=2)
    # A7：英文
    fix(doc, "I hereby give my consent my personal data processing by BorsodChem Zrt. as follows:",
        "I hereby consent to the processing of my personal data by BorsodChem Zrt. for the following purposes:")
    fix(doc, "family union – and other visa types applications,", "family reunification and other visa applications,")
    fix(doc, "making renting contracts,", "making rental contracts,")
    fix(doc, "family child allowance and tax allowance application,", "family allowance and tax allowance applications,")
    # C：CARINVISA 是笔误（2026-10-31 确认）
    fix(doc, "that CARINVISA shares my sensitive data", "that BorsodChem Zrt. shares my sensitive data")
    check_done(doc, ["CARINVISA"])
    doc.save(OUT / "gdpr_declarations.docx")


if __name__ == "__main__":
    build_application()
    build_order()
    build_gdpr()
    print("已生成 taj_application、taj_order、gdpr_declarations")
