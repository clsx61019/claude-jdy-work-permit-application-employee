"""
一次性工具：把简道云的四份模板转换成本项目的模板（规则见 docs/rules/住宿证明雇主通知员工声明.md）。

- Accommodation Certificate → accommodation_certificate.docx（住宿证明，Polimer 经营的住处出具）
- Employer Notification → employer_notification.docx（雇主用工情况报告）
- Declarations（三页）→ 拆成三份：
    request_previous_documents.docx（沿用材料请求）
    request_delivery_address.docx（通信和寄卡地址请求）
    declaration_address_unchanged.docx（住址未变声明）
- Declaration for employer → employee_consent.docx（员工同意雇主代为递交）

双语的都改成左栏匈牙利语、右栏英语。

用法: python tools/build_employee_documents.py
"""
import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from service.docx_text import para_text, paragraphs, replace_in_paragraph  # noqa: E402
from tools.bilingual import hungarian_left  # noqa: E402

W = r"\$\{[^}]*?#_widget_"  # 占位符开头
OUT = ROOT / "templates/hu"


def fix(doc, old, new, regex=False, expect=1):
    pattern = old if regex else re.escape(old)
    n = sum(replace_in_paragraph(p, pattern, lambda m: new) for p in paragraphs(doc))
    if n != expect:
        raise ValueError(f"「{old[:50]}…」应命中 {expect} 处，实际 {n} 处")


def check_done(doc):
    left = [para_text(p) for p in paragraphs(doc) if "${" in para_text(p)]
    if left:
        raise ValueError(f"还有没换掉的字段：{left}")


def rows(doc):
    return [tr for tbl in doc.element.body.iter(qn("w:tbl")) for tr in tbl.findall(qn("w:tr"))]


def remove_empty_rows(doc, indexes):
    """按序号删除第一个表格里的空行（模板用空行撑间距，签字栏被挤到第二页）"""
    trs = doc.tables[0]._tbl.findall(qn("w:tr"))
    for i in sorted(indexes, reverse=True):
        if para_text(trs[i]).strip():
            raise ValueError(f"第 {i} 行不是空行：{para_text(trs[i])[:30]}")
        trs[i].getparent().remove(trs[i])


def put(tc, text, like):
    """在空单元格里写文字，格式照抄 like 段落里的第一个 run"""
    p = tc.find(qn("w:p"))
    run = copy.deepcopy(next(like.iter(qn("w:r"))))
    for t in run.findall(qn("w:t"))[1:]:
        run.remove(t)
    run.find(qn("w:t")).text = text
    p.append(run)


def remove_highlight(doc):
    for h in list(doc.element.body.iter(qn("w:highlight"))):
        h.getparent().remove(h)


# ---------- 员工签的几份（双语，原来已是左匈右英） ----------

def employee_header(doc):
    """A1、A3：姓名、出生地（按出生国）、出生日期"""
    fix(doc, W + r"1728475331388\}", "{{name}}", regex=True, expect=2)
    fix(doc, r"szül\. hely: " + W + r"1732527942947\}, " + W + r"1732527942838\}; szül\.idő: " + W + r"1773645284571\}",
        "szül. hely: {{birth_country_hu}}, {{birth_place}}; szül. idő: {{birth_date_hu}}", regex=True)
    fix(doc, r"place of birth: China, " + W + r"1732527942838\}; date of birth: " + W + r"1773645284570\}",
        "place of birth: {{birth_country_en}}, {{birth_place}}; date of birth: {{birth_date_en}}", regex=True)


def employee_footer(doc):
    """A2：日期两栏原来对调（匈语栏印英语日期）；B9：签字地点按递交地点"""
    fix(doc, "Kazincbarcika, " + "${辅助_签字日期1（英语）#_widget_1773645284867}", "{{place_hu}}, {{sign_date_hu}}")
    fix(doc, "Kazincbarcika, " + "${辅助_签字日期1（匈牙利语）#_widget_1773645284868}", "{{place_en}}, {{sign_date_en}}")


def keep_page(src, start):
    """Declarations 原模板三页靠空段隔开：每页是「正文表格 + 签字表格」。只留第 start 页"""
    doc = Document(src)
    body = doc.element.body
    tables = body.findall(qn("w:tbl"))
    keep = tables[start * 2: start * 2 + 2]
    for el in list(body):
        if el.tag == qn("w:sectPr") or el in keep:
            continue
        body.remove(el)
    # 正文和签字栏之间留三段空行（原模板如此），签字栏后留一段
    for _ in range(3):
        keep[0].addnext(copy.deepcopy(_empty_paragraph()))
    keep[1].addnext(_empty_paragraph())
    return doc


def _empty_paragraph():
    from docx.oxml import OxmlElement
    return OxmlElement("w:p")


def build_requests():
    src = ROOT / "templates/jdy/declarations_jdy_original.docx"

    doc = keep_page(src, 0)
    employee_header(doc)
    employee_footer(doc)
    remove_highlight(doc)
    check_done(doc)
    doc.save(OUT / "request_previous_documents.docx")

    doc = keep_page(src, 1)
    employee_header(doc)
    employee_footer(doc)
    # A10：邮箱、地址、收件人取公司配置；英文改通顺
    fix(doc, "tájékoztató leveleiket a renata.simon@borsodchem.eu e-mail címre és elkészült okmányomat a "
             "3700 Kazincbarcika, Bolyai tér 1. szám alá, Simon Renáta névre",
        "tájékoztató leveleiket {{contact_email_with_article}} e-mail címre és elkészült okmányomat "
        "{{delivery_address_with_article}} szám alá, {{contact_name_hu}} névre")
    fix(doc, r"hereby kindly ask the authority to send all related correspondence.*$",
        "hereby kindly ask the authority to send all correspondence related to my residence permit application "
        "to the e-mail address {{contact_email}} and the issued residence card to {{delivery_address}}, "
        "to the attention of {{contact_name_en}}, for the sake of easier administration.", regex=True)
    remove_highlight(doc)
    check_done(doc)
    doc.save(OUT / "request_delivery_address.docx")

    doc = keep_page(src, 2)
    employee_header(doc)
    employee_footer(doc)
    addr = (W + r"1728549508506\} " + W + r"1728549508516\}, " + W + r"1728549508519\} " + W + r"1728549508520\} "
            + W + r"1728549508521\}")
    # A5、A11：地址统一写法，补句点，冠词按读音
    fix(doc, "tartózkodási helyem továbbra is a " + addr, "tartózkodási helyem továbbra is {{hu_address_with_article}}.",
        regex=True)
    fix(doc, "hereby declare that my HU address is unchanged " + addr,
        "hereby declare that my place of residence in Hungary remains {{hu_address}}.", regex=True)
    # 这一页是声明，不是请求
    fix(doc, r"^Kérelem $", "Nyilatkozat", regex=True)
    fix(doc, r"^Request$", "Declaration", regex=True)
    remove_highlight(doc)
    check_done(doc)
    doc.save(OUT / "declaration_address_unchanged.docx")


def build_consent():
    doc = Document(ROOT / "templates/jdy/employee_consent_jdy_original.docx")
    employee_header(doc)
    employee_footer(doc)
    # B8：居留许可类型按字段
    fix(doc, "hogy Vállalaton-belüli áthelyezésre irányuló tartózkodási engedély iránti kérelmemet munkáltatóm "
             "kezdeményezze az EnterHungary felületén.",
        "hogy {{permit_hu}} iránti kérelmemet munkáltatóm kezdeményezze az Enter Hungary felületén.")
    fix(doc, "hereby give my consent to my employer to submit my request on EnterHungary for Intra-company Transfer "
             "residence permit application.",
        "hereby give my consent to my employer submitting my application for {{permit_en}} via Enter Hungary.")
    remove_highlight(doc)
    check_done(doc)
    doc.save(OUT / "employee_consent.docx")


# ---------- 住宿证明（Polimer 信头，原来英语在左） ----------

def build_accommodation():
    doc = Document(ROOT / "templates/jdy/accommodation_certificate_jdy_original.docx")
    fix(doc, r"^We hereby confirm that .*$",
        "We hereby confirm that {{name}} (date of birth: {{birth_date_en}}), {{nationality_en}} citizen, will live in "
        "our hotel, {{hotel}} (address: {{hu_address}}), during {{his_her}} stay in Hungary.", regex=True)
    # A6：匈语语序
    fix(doc, r"^Ezúton igazoljuk, hogy .*$",
        "Ezúton igazoljuk, hogy {{name}} (születési dátum: {{birth_date_hu}}) {{nationality_hu}} állampolgár "
        "magyarországi tartózkodása alatt {{hotel_with_article}} szállodánkban (cím: {{hu_address}}) fog lakni.",
        regex=True)
    fix(doc, "Kazincbarcika, " + "${REF_Date of Signatures 1_EN#_widget_1773645284867}", "Kazincbarcika, {{sign_date_en}}")
    fix(doc, "Kazincbarcika, " + "${REF_Date of Signatures 1_HU#_widget_1773645284868}", "Kazincbarcika, {{sign_date_hu}}")
    # A4：签字栏被挤到第二页，删掉几行撑间距的空行
    remove_empty_rows(doc, [7, 8, 12, 13, 14, 18, 19, 20, 21])
    hungarian_left(doc)
    check_done(doc)
    doc.save(OUT / "accommodation_certificate.docx")


# ---------- 雇主通知（匈语） ----------

def build_notification():
    doc = Document(ROOT / "templates/jdy/employer_notification_jdy_original.docx")
    fix(doc, W + r"1732604018708\}", "{{employer}}", regex=True)
    # A9：联系方式、地址、税号、注册号取配置
    fix(doc, "+36 30 960 4086, renata.simon@borsodchem.eu", "{{contact}}")
    fix(doc, W + r"1732604018771\}", "{{employer_address}}", regex=True)
    fix(doc, W + r"1732604018791\}", "{{employer_tax_no}}", regex=True)
    fix(doc, W + r"1733392328951\}", "{{employer_registration_no}}", regex=True)
    fix(doc, "Foglakoztatott neve:", "Foglalkoztatott neve:")  # A7
    fix(doc, W + r"1728475331388\}", "{{name}}", regex=True)
    fix(doc, r"^" + W + r"1732527942838\}, .*$", "{{birth}}", regex=True)
    fix(doc, r"^kínai$", "{{nationality_hu}}", regex=True)  # A8
    trs = doc.tables[0]._tbl.findall(qn("w:tr"))
    like = trs[4].findall(qn("w:tc"))[0].find(qn("w:p"))
    # B4：居留卡号；a) 的日期（有「劳动关系实际起始日期」时由代码填，b)、c) 两行删掉）
    card = next(tr for tr in trs if para_text(tr).startswith("Tartózkodási engedélyének"))
    put(card.findall(qn("w:tc"))[1], "{{card_no}}", like)
    start = next(tr for tr in trs if para_text(tr).startswith("a) foglalkoztatását"))
    put(start.findall(qn("w:tc"))[1], "{{start_date}}", like)
    # B5：落款日期
    fix(doc, "Kazincbarcika, ......... év ......... hó ......... nap", "Kazincbarcika, {{today}}")
    # B6：签字人印雇主代表
    fix(doc, "Cégvezető (cégszerű aláírása)", "{{representative}}")
    sig = next(p for p in paragraphs(doc) if para_text(p) == "{{representative}}")
    sub = copy.deepcopy(sig)
    sig.addnext(sub)
    replace_in_paragraph(sub, re.escape("{{representative}}"), lambda m: "(cégszerű aláírás)")
    # A4：签字栏被挤到第二页
    remove_empty_rows(doc, [10, 11, 22, 23, 27, 28])
    check_done(doc)
    doc.save(OUT / "employer_notification.docx")


if __name__ == "__main__":
    build_accommodation()
    build_notification()
    build_requests()
    build_consent()
    print("已生成 accommodation_certificate、employer_notification、request_previous_documents、"
          "request_delivery_address、declaration_address_unchanged、employee_consent")
